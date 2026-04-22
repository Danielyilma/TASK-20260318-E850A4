from __future__ import annotations

import csv
import io
import math
import uuid
from datetime import datetime
from pathlib import Path

from fastapi import HTTPException, status
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.time import utcnow
from app.models.audit_log import AuditLog
from app.models.funding_account import FundingAccount
from app.models.generated_report import GeneratedReport
from app.models.registration import Registration
from app.models.transaction_record import TransactionRecord
from app.models.user import User


class Phase5ReportsService:
    def __init__(self, db: Session) -> None:
        self.db = db

    def _reports_dir(self) -> Path:
        root = Path(get_settings().upload_root) / "reports"
        root.mkdir(parents=True, exist_ok=True)
        return root

    def generate_reconciliation(
        self,
        user: User,
        *,
        activity_id: uuid.UUID | None,
        start_date: datetime | None,
        end_date: datetime | None,
        file_format: str,
    ) -> GeneratedReport:
        file_format = file_format.lower()
        if file_format not in ("xlsx", "csv", "pdf"):
            from app.core.http_errors import bad_request

            raise bad_request("VALIDATION_ERROR", "format must be one of xlsx, csv, pdf")

        q = (
            select(TransactionRecord, FundingAccount, Registration)
            .join(FundingAccount, TransactionRecord.funding_account_id == FundingAccount.id)
            .join(Registration, FundingAccount.registration_id == Registration.id)
        )
        filters = []
        if activity_id is not None:
            filters.append(Registration.activity_id == activity_id)
        if start_date is not None:
            filters.append(TransactionRecord.created_at >= start_date)
        if end_date is not None:
            filters.append(TransactionRecord.created_at <= end_date)
        if filters:
            q = q.where(*filters)
        rows = list(self.db.execute(q).all())

        rid = uuid.uuid4()
        fname = f"reconciliation_{rid.hex[:8]}.{file_format}"
        path = self._reports_dir() / f"{rid}.{file_format}"

        if file_format == "csv":
            buf = io.StringIO()
            w = csv.writer(buf)
            w.writerow(["transaction_id", "type", "amount", "category", "registration_id", "created_at"])
            for tx, fa, reg in rows:
                w.writerow([str(tx.id), tx.type, float(tx.amount), tx.category, str(reg.id), tx.created_at.isoformat()])
            raw = buf.getvalue().encode("utf-8")
            path.write_bytes(raw)
        elif file_format == "xlsx":
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "transactions"
            ws.append(["transaction_id", "type", "amount", "category", "registration_id", "created_at"])
            for tx, fa, reg in rows:
                ws.append([str(tx.id), str(tx.type), float(tx.amount), tx.category, str(reg.id), tx.created_at.isoformat()])
            wb.save(path)
        else:
            from reportlab.pdfgen import canvas as rl_canvas

            c = rl_canvas.Canvas(str(path), pagesize=letter)
            y = 750
            c.drawString(50, y, "Reconciliation report")
            y -= 24
            for tx, fa, reg in rows[:80]:
                line = f"{tx.id} {tx.type} {tx.amount} {tx.category} reg={reg.id}"
                c.drawString(50, y, line[:100])
                y -= 14
                if y < 50:
                    c.showPage()
                    y = 750
            c.save()

        size = path.stat().st_size
        rep = GeneratedReport(
            id=rid,
            report_type="reconciliation",
            status="completed",
            file_path=str(path),
            file_name=fname,
            file_size_bytes=size,
            file_format=file_format,
            created_by=user.id,
            activity_id=activity_id,
            created_at=utcnow(),
        )
        self.db.add(rep)
        self.db.commit()
        self.db.refresh(rep)
        return rep

    def generate_audit_report(
        self,
        user: User,
        *,
        start_date: datetime | None,
        end_date: datetime | None,
        user_id: uuid.UUID | None,
        action_filter: str | None,
        file_format: str,
    ) -> GeneratedReport:
        file_format = file_format.lower()
        if file_format not in ("xlsx", "csv", "pdf"):
            from app.core.http_errors import bad_request

            raise bad_request("VALIDATION_ERROR", "format must be one of xlsx, csv, pdf")
        q = select(AuditLog)
        filters = []
        if start_date:
            filters.append(AuditLog.created_at >= start_date)
        if end_date:
            filters.append(AuditLog.created_at <= end_date)
        if user_id:
            filters.append(AuditLog.user_id == user_id)
        if action_filter:
            filters.append(AuditLog.action == action_filter)
        if filters:
            q = q.where(*filters)
        logs = list(self.db.scalars(q.order_by(AuditLog.created_at.desc()).limit(5000)).all())

        rid = uuid.uuid4()
        fname = f"audit_{rid.hex[:8]}.{file_format}"
        path = self._reports_dir() / f"{rid}.{file_format}"

        if file_format == "csv":
            buf = io.StringIO()
            w = csv.writer(buf)
            w.writerow(["id", "user_id", "username", "action", "resource_type", "resource_id", "created_at"])
            for a in logs:
                w.writerow(
                    [
                        str(a.id),
                        str(a.user_id) if a.user_id else "",
                        a.username or "",
                        a.action,
                        a.resource_type or "",
                        str(a.resource_id) if a.resource_id else "",
                        a.created_at.isoformat(),
                    ]
                )
            path.write_bytes(buf.getvalue().encode("utf-8"))
        elif file_format == "xlsx":
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.append(["id", "user_id", "username", "action", "resource_type", "resource_id", "created_at"])
            for a in logs:
                ws.append(
                    [
                        str(a.id),
                        str(a.user_id) if a.user_id else "",
                        a.username or "",
                        a.action,
                        a.resource_type or "",
                        str(a.resource_id) if a.resource_id else "",
                        a.created_at.isoformat(),
                    ]
                )
            wb.save(path)
        else:
            from reportlab.pdfgen import canvas as rl_canvas

            c = rl_canvas.Canvas(str(path), pagesize=letter)
            y = 750
            c.drawString(50, y, "Audit trail report")
            y -= 30
            for a in logs[:100]:
                c.drawString(50, y, f"{a.created_at} {a.action} {a.resource_type} {a.resource_id}")
                y -= 16
                if y < 40:
                    c.showPage()
                    y = 750
            c.save()

        size = path.stat().st_size
        rep = GeneratedReport(
            id=rid,
            report_type="audit",
            status="completed",
            file_path=str(path),
            file_name=fname,
            file_size_bytes=size,
            file_format=file_format,
            created_by=user.id,
            activity_id=None,
            created_at=utcnow(),
        )
        self.db.add(rep)
        self.db.commit()
        self.db.refresh(rep)
        return rep

    def generate_compliance(self, user: User, *, activity_id: uuid.UUID, file_format: str) -> GeneratedReport:
        file_format = file_format.lower()
        if file_format not in ("xlsx", "csv", "pdf"):
            from app.core.http_errors import bad_request

            raise bad_request("VALIDATION_ERROR", "format must be one of xlsx, csv, pdf")

        regs = list(
            self.db.scalars(select(Registration).where(Registration.activity_id == activity_id)).all()
        )
        rid = uuid.uuid4()
        fname = f"compliance_{rid.hex[:8]}.{file_format}"
        path = self._reports_dir() / f"{rid}.{file_format}"

        if file_format == "csv":
            buf = io.StringIO()
            w = csv.writer(buf)
            w.writerow(["registration_id", "status", "requested_funding"])
            for r in regs:
                w.writerow([str(r.id), r.status.value if hasattr(r.status, "value") else r.status, float(r.requested_funding)])
            path.write_bytes(buf.getvalue().encode("utf-8"))
        elif file_format == "xlsx":
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.append(["registration_id", "status", "requested_funding"])
            for r in regs:
                st = r.status.value if hasattr(r.status, "value") else str(r.status)
                ws.append([str(r.id), st, float(r.requested_funding)])
            wb.save(path)
        else:
            from reportlab.pdfgen import canvas as rl_canvas

            c = rl_canvas.Canvas(str(path), pagesize=letter)
            y = 750
            c.drawString(50, y, f"Compliance summary activity={activity_id}")
            y -= 24
            for r in regs[:80]:
                st = r.status.value if hasattr(r.status, "value") else str(r.status)
                c.drawString(50, y, f"{r.id} {st}")
                y -= 14
                if y < 40:
                    c.showPage()
                    y = 750
            c.save()

        size = path.stat().st_size
        rep = GeneratedReport(
            id=rid,
            report_type="compliance",
            status="completed",
            file_path=str(path),
            file_name=fname,
            file_size_bytes=size,
            file_format=file_format,
            created_by=user.id,
            activity_id=activity_id,
            created_at=utcnow(),
        )
        self.db.add(rep)
        self.db.commit()
        self.db.refresh(rep)
        return rep

    def list_reports(
        self,
        *,
        page: int,
        per_page: int,
        report_type: str | None,
        file_format: str | None,
    ) -> dict:
        q = select(GeneratedReport)
        filters = []
        if report_type:
            filters.append(GeneratedReport.report_type == report_type)
        if file_format:
            filters.append(GeneratedReport.file_format == file_format)
        if filters:
            q = q.where(*filters)
        total = int(self.db.scalar(select(func.count()).select_from(q.subquery())) or 0)
        rows = self.db.scalars(
            q.order_by(GeneratedReport.created_at.desc()).offset((page - 1) * per_page).limit(per_page)
        ).all()
        pages = math.ceil(total / per_page) if total else 0
        items = [
            {
                "report_id": str(r.id),
                "report_type": r.report_type,
                "status": r.status,
                "file_name": r.file_name,
                "file_size_bytes": r.file_size_bytes,
                "format": r.file_format,
                "created_at": r.created_at.isoformat().replace("+00:00", "Z"),
            }
            for r in rows
        ]
        return {"items": items, "total": total, "page": page, "per_page": per_page, "pages": pages}

    def download_path(self, report_id: uuid.UUID) -> Path:
        r = self.db.get(GeneratedReport, report_id)
        if r is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Report not found")
        p = Path(r.file_path)
        if not p.is_file():
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Report not found")
        return p
